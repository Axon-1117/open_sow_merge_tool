from types import SimpleNamespace

from openpyxl import load_workbook

from sow_merge_tool import legacy_core as core
from sow_merge_tool.branch_submit import BranchSubmitEngine
from sow_merge_tool.svn_status_provider import SvnStatusRecord


def test_new_source_has_no_base_but_modified_source_requires_it(tmp_path):
    engine = BranchSubmitEngine(str(tmp_path), require_remote_freshness=True)
    for status in ('added', 'unversioned', 'modified'):
        record = SvnStatusRecord(str(tmp_path / 'new.xlsx'), node_status=status,
                                 repository_revision=10)
        evidence = engine._freshness_evidence(branch='develop', path=record.path,
                                             record=record, expected_revision=None,
                                             phase='source')
        assert evidence['state'] == ('unknown' if status == 'modified' else 'not_applicable')


def test_literal_does_not_build_formula_mapping():
    class UnexpectedMapping:
        def items(self):
            raise AssertionError('Literal cell must not enumerate formula mappings')
    assert core._canonicalize_formula_column_references('plain text', UnexpectedMapping()) == 'plain text'


def test_complete_read_cache_preserves_padding_and_does_not_reparse():
    ws = SimpleNamespace(_sow_complete_read_rows=((1, '中文'), (2, '=A1')))
    assert core._read_rows_into_cache(ws, [1, 2], 3) == {
        1: (1, '中文', None), 2: (2, '=A1', None),
    }


def test_screen_reports_formula_presence(tmp_path):
    from openpyxl import Workbook

    from sow_merge_tool.sheet_screening import sheet_fingerprints
    path = tmp_path / 'book.xlsx'
    wb = Workbook()
    wb.active['A1'] = '=1+2'
    wb.create_sheet('literal')['A1'] = 3
    wb.save(path)
    wb.close()
    flags = {}
    sheet_fingerprints(str(path), formula_flags=flags)
    assert flags == {'Sheet': True, 'literal': False}
    wb = load_workbook(path, read_only=True)
    assert wb['Sheet']['A1'].value == '=1+2'
    wb.close()
