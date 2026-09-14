from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sow_merge_tool.legacy_core import (
    _row_sig_list_for_ws,
    _workbook_health_evidence,
    _worksheet_scan_bounds,
)


def _write_book(path: Path, values: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row_index, row in enumerate(values, start=1):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column_index, value=value)
    workbook.save(path)
    workbook.close()


def _force_a1_dimension(path: Path) -> None:
    rewritten = path.with_name(f"{path.stem}.a1{path.suffix}")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                start = payload.find(b"<dimension")
                end = payload.find(b"/>", start)
                payload = payload[:start] + b'<dimension ref="A1"/>' + payload[end + 2:]
            target.writestr(info, payload)
    rewritten.replace(path)


def test_read_only_bounds_recover_rows_from_misleading_a1_dimension(tmp_path):
    base = tmp_path / "base.xlsx"
    mine = tmp_path / "mine.xlsx"
    values = [["id", "value"], [1, "same"], [2, "base"]]
    _write_book(base, values)
    _write_book(mine, [["id", "value"], [1, "same"], [2, "mine"]])
    _force_a1_dimension(base)
    _force_a1_dimension(mine)
    health = _workbook_health_evidence(str(mine))
    assert health["ready"] is True
    assert health["worksheet_count"] == 1
    assert health["a1_dimension_parts"] == ("xl/worksheets/sheet1.xml",)

    base_wb = load_workbook(base, read_only=True, data_only=False)
    mine_wb = load_workbook(mine, read_only=True, data_only=False)
    try:
        base_ws = base_wb.active
        mine_ws = mine_wb.active
        assert _worksheet_scan_bounds(base_ws) == (3, 2)
        assert _worksheet_scan_bounds(mine_ws) == (3, 2)
        assert _row_sig_list_for_ws(base_ws, 3, 2) != _row_sig_list_for_ws(mine_ws, 3, 2)
    finally:
        base_wb.close()
        mine_wb.close()
