"""Regression: premerge caches survive preview, no-edit save, and sheet copying."""
import os
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
import sow_merge_tool as sm
from _smoke_test_column_native_save_replay import _fake_app


def make_book(path, changed=False, extra=False, missing=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["id", "value", "formula"])
    for i in range(1, 305):
        ws.append([i, "new" if changed and i == 1 else "old", "=2+2" if changed and i == 2 else "=1+1"])
    if extra:
        wb.create_sheet("Added")["A1"] = '=IF(1,"","")'
    wb.save(path)
    wb.close()
    ops = {("Data", i + 1, 3): "=2+2" if changed and i == 2 else "=1+1" for i in range(1, 305)}
    values = {key: 4 if changed and key[1] == 3 else 2 for key in ops}
    # Cache types that must survive without coercion, including valid empty string.
    for row, value in ((4, 0), (5, False), (6, "text"), (7, ""), (8, "#N/A")):
        values[("Data", row, 3)] = value
    if extra:
        ops[("Added", 1, 1)] = '=IF(1,"","")'
        values[("Added", 1, 1)] = ""
    if missing:
        values[("Data", 305, 3)] = None
    sm._build_manual_merge_xlsx_via_zip(path, path + ".patched", ops, cached_values=values)
    os.replace(path + ".patched", path)


def caches(path):
    q = lambda name: "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}" + name
    with zipfile.ZipFile(path) as z:
        payloads = {name: z.read(name) for name in z.namelist()}
    result = {}
    for sheet, part in sm._ooxml_sheet_part_map(payloads).items():
        for cell in ET.fromstring(payloads[part]).iter(q("c")):
            if cell.find(q("f")) is not None:
                value = cell.find(q("v"))
                result[(sheet, cell.get("r"))] = (cell.get("t"), value is not None, None if value is None else value.text)
    return result


def main():
    with tempfile.TemporaryDirectory(prefix="sow-premerge-cache-") as root:
        base, mine, theirs, target = [os.path.join(root, x + ".xlsx") for x in ("base", "mine", "theirs", "target")]
        make_book(base)
        make_book(mine)
        make_book(theirs)
        original = open(mine, "rb").read()
        for changed, extra in ((False, False), (True, False), (True, True)):
            make_book(theirs, changed=changed, extra=extra)
            conflicts, preview, _ = sm._merge_three_way(base, mine, theirs, target, save_merged=False)
            try:
                assert not conflicts, conflicts
                assert caches(preview) == caches(theirs), (changed, extra)
                sm._validate_formula_caches_for_save(preview)
                if not changed:
                    assert open(preview, "rb").read() == original
                wb = load_workbook(preview)
                assert wb["Data"]["C3"].value == ("=2+2" if changed else "=1+1")
                assert wb["Data"]["B2"].value == ("new" if changed else "old")
                wb.close()
            finally:
                os.remove(preview)
        # Direct publication validates before replacing the user's file.
        make_book(theirs, missing=True)
        make_book(mine, missing=True)
        make_book(base, missing=True)
        with open(target, "wb") as f:
            f.write(original)
        try:
            sm._merge_three_way(base, mine, theirs, target, save_merged=True)
        except RuntimeError as exc:
            assert "Data!C305" in str(exc), str(exc)
        else:
            raise AssertionError("missing cache published")
        assert open(target, "rb").read() == original
        # Public merged-save entry must retain user state and never resolve/exit.
        app = _fake_app(mine, theirs)
        app.merged_path = target
        app.merge_mode = True
        app.has_base = True
        app.initial_conflict_cell_count = 0
        app.sheet_level_conflicts = []
        app.modified_a = True
        app.manual_a_cell_ops[("Data", 2, 2)] = "pending"
        app._wb_a_edit = load_workbook(mine)
        app._guard_save_readiness = lambda *a: True
        app._ensure_edit_loaded = lambda: None
        app.build_manual_merge_output_file = lambda: sm._create_startup_candidate_copy(mine, "test")
        app._begin_interactive_action = lambda: None
        app._end_interactive_action = lambda: None
        app._with_progress = lambda title, text, task, **kw: task(lambda *a: None)
        with patch.object(sm.messagebox, "askyesno", return_value=True), patch.object(sm.messagebox, "showerror") as error, patch.object(sm.messagebox, "showinfo") as info, patch.object(sm, "_try_svn_resolve") as resolve:
            app.save_merged_and_exit()
            assert error.called and "Data!C305" in error.call_args.args[1]
            assert not info.called and not resolve.called
        assert app.modified_a and app.manual_a_cell_ops
        assert open(target, "rb").read() == original
        # Retry with complete caches: no-edit final save keeps all 304 results.
        make_book(mine)
        app.manual_a_cell_ops.clear()
        app._shutdown_root = lambda: None
        with patch.object(sm.messagebox, "askyesno", return_value=True), patch.object(sm.messagebox, "showerror") as error, patch.object(sm.messagebox, "showinfo"), patch.object(sm, "_has_svn_conflict_artifacts", return_value=False):
            try:
                app.save_merged_and_exit()
            except SystemExit as exc:
                assert exc.code == 0
            else:
                raise AssertionError("valid save did not complete")
            assert not error.called, error.call_args
        assert caches(target) == caches(mine)
        assert not app.modified_a
        app._wb_a_edit.close()
    print("PASS: 304 caches, typed/blank results, adopted formula, new Sheet, and blocked publication")


if __name__ == "__main__":
    main()
