"""Regression for unequal independent tail appends on a declared-key Sheet."""

from __future__ import annotations

import os
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

import sow_merge_tool as sm


_SHEET = "ConditionData@design"
_HEADER = [
    [
        "id@rely_id#key#ConditionData",
        "key@index_uniq_cached@constid",
        "type@ref_Condition",
        "param1",
    ],
    ["uint32", "string", "@ref", "string"],
    ["generated id", "condition key", "condition type", "parameter"],
    [None, None, None, None],
]


def _write(path: str, rows: list[list[object]], sheet_name: str = _SHEET) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _snapshot(path: str, side: str, sheet_name: str = _SHEET):
    return sm._stream_selected_sheet_snapshot(path, path, sheet_name, side)


def _wait_for_ready(app, timeout: float = 20.0):
    deadline = time.monotonic() + timeout
    view = None
    while time.monotonic() < deadline:
        app._request_edit_preload()
        app.root.update_idletasks()
        app.root.update()
        view = app.sheet_views.get(_SHEET)
        if (
            view is not None
            and view._data_ready
            and app._is_sheet_exact_current(_SHEET)
            and app._edit_workbooks_ready()
            and view._derive_lifecycle_state() == "READY"
        ):
            return view
        time.sleep(0.01)
    raise AssertionError((app._sheet_exact_entry(_SHEET), view))


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="sow_declared_appends_3way_") as root:
        base_path = os.path.join(root, "base.xlsx")
        mine_path = os.path.join(root, "mine.xlsx")
        theirs_path = os.path.join(root, "theirs.xlsx")
        common = [[None, "quest_complete_402001", "quest_complete", 402001]]
        mine_additions = [
            [None, "grid_discover_day_2", "server_open_day", 2],
            [None, "grid_discover_day_3", "server_open_day", 3],
        ]
        theirs_additions = [
            [None, f"recharge_{condition}", "recharge", condition]
            for condition in range(30102, 30107)
        ]
        _write(base_path, _HEADER + common)
        # Match the real workbook shape: Mine has a blank separator before its
        # two records, while Theirs appends five records directly.
        _write(
            mine_path,
            _HEADER + common + [[None, None, None, None]] + mine_additions,
        )
        _write(theirs_path, _HEADER + common + theirs_additions)

        base = _snapshot(base_path, "BASE")
        mine = _snapshot(mine_path, "A")
        theirs = _snapshot(theirs_path, "B")
        pairwise = sm._align_selected_sheet_snapshots(mine, theirs)
        # Unique declared keys must prevent the pairwise display alignment from
        # pairing unrelated additions merely because they share one tail gap.
        assert sum(
            mine_row is None and theirs_row is not None
            for mine_row, theirs_row in pairwise.row_pairs
        ) == 5, pairwise.row_pairs
        assert sum(
            mine_row is not None and theirs_row is None
            for mine_row, theirs_row in pairwise.row_pairs
        ) == 3, pairwise.row_pairs

        result = sm._compare_selected_sheet_snapshots(mine, theirs, base)
        expected_tail = (
            (6, None),
            (7, None),
            (8, None),
            (None, 6),
            (None, 7),
            (None, 8),
            (None, 9),
            (None, 10),
        )
        assert not result.unresolved
        assert result.row_pairs[-8:] == expected_tail, result.row_pairs
        assert result.base_rows_by_pair[-8:] == (None,) * 8
        assert all(
            result.pair_diff_cols[index] == frozenset((-1,))
            for index in range(len(result.row_pairs) - 8, len(result.row_pairs))
        )
        assert not any(
            mine_row is not None
            and theirs_row is not None
            and result.base_rows_by_pair[index] is None
            for index, (mine_row, theirs_row) in enumerate(result.row_pairs)
        )

        cache = sm._snapshot_result_to_sheet_cache_immutable(
            _SHEET, result, mine, theirs, base, has_base=True,
        )
        assert cache["prepared_complete"]
        assert tuple(cache["row_pairs"][-8:]) == expected_tail

        # Reopening after the five incoming records have become the new Base
        # must not let Mine's unrelated tail rows claim those Base coordinates.
        reopen_base_path = os.path.join(root, "reopen_base.xlsx")
        reopen_mine_path = os.path.join(root, "reopen_mine.xlsx")
        reopen_theirs_path = os.path.join(root, "reopen_theirs.xlsx")
        next_addition = [None, "city_event_reward_3", "city_event_reward", 3]
        _write(reopen_base_path, _HEADER + common + theirs_additions)
        _write(
            reopen_mine_path,
            _HEADER + common + [[None, None, None, None]] + mine_additions,
        )
        _write(
            reopen_theirs_path,
            _HEADER + common + theirs_additions
            + [[None, None, None, None], next_addition],
        )
        reopen_base = _snapshot(reopen_base_path, "BASE")
        reopen_mine = _snapshot(reopen_mine_path, "A")
        reopen_theirs = _snapshot(reopen_theirs_path, "B")
        reopen_result = sm._compare_selected_sheet_snapshots(
            reopen_mine, reopen_theirs, reopen_base,
        )
        assert not reopen_result.unresolved
        assert reopen_result.row_pairs[-10:] == (
            (6, None), (7, None), (8, None),
            (None, 6), (None, 7), (None, 8), (None, 9), (None, 10),
            (None, 11), (None, 12),
        ), reopen_result.row_pairs
        assert reopen_result.base_rows_by_pair[-10:] == (
            None, None, None, 6, 7, 8, 9, 10, None, None,
        )
        assert sm._snapshot_result_base_row_mappings(
            reopen_result, reopen_mine, reopen_theirs, reopen_base,
        ) is not None
        reopen_cache = sm._snapshot_result_to_sheet_cache_immutable(
            _SHEET, reopen_result, reopen_mine, reopen_theirs, reopen_base,
            has_base=True,
        )
        assert reopen_cache["prepared_complete"]

        # An identical unkeyed Sheet can be ambiguous to the general content
        # matcher (duplicate row identities), but exact selected-Sheet equality
        # proves the physical 1:1 result without relying on workbook metadata.
        unkeyed_sheet = "TCondition@design"
        unkeyed_rows = [
            ["ID", "Type"], ["int", "int"], [1, 5], [1, 5],
        ]
        unkeyed_paths = [
            os.path.join(root, f"unkeyed_{side}.xlsx")
            for side in ("base", "mine", "theirs")
        ]
        for path in unkeyed_paths:
            _write(path, unkeyed_rows, unkeyed_sheet)
        unkeyed_base = _snapshot(unkeyed_paths[0], "BASE", unkeyed_sheet)
        unkeyed_mine = _snapshot(unkeyed_paths[1], "A", unkeyed_sheet)
        unkeyed_theirs = _snapshot(unkeyed_paths[2], "B", unkeyed_sheet)
        assert sm._align_selected_sheet_snapshots(
            unkeyed_mine, unkeyed_theirs,
        ).unresolved
        unkeyed_result = sm._compare_selected_sheet_snapshots(
            unkeyed_mine, unkeyed_theirs, unkeyed_base,
        )
        assert not unkeyed_result.unresolved
        assert unkeyed_result.physical_identity
        unkeyed_cache = sm._snapshot_result_to_sheet_cache_immutable(
            unkeyed_sheet, unkeyed_result,
            unkeyed_mine, unkeyed_theirs, unkeyed_base,
            has_base=True,
        )
        assert unkeyed_cache["prepared_complete"]
        original_scheduler = sm.SowMergeApp._schedule_formula_cache_prompt
        sm.SowMergeApp._schedule_formula_cache_prompt = lambda _self: None
        app = None
        output = None
        try:
            app = sm.SowMergeApp(
                mine_path,
                theirs_path,
                merge_mode=True,
                merged_path=os.path.join(root, "merged.xlsx"),
                base_path=base_path,
                initial_sheet=_SHEET,
            )
            view = _wait_for_ready(app)
            remote_keys = [row[1] for row in theirs_additions]
            detected_keys = [
                app.ws_b_val(_SHEET).cell(row=theirs_row, column=2).value
                for mine_row, theirs_row in view.row_pairs
                if mine_row is None and theirs_row is not None
            ]
            assert detected_keys == remote_keys, detected_keys

            # Every Theirs row must remain directly actionable after earlier
            # insertions rebuild the row model.
            for key in remote_keys:
                pair_index = next(
                    index
                    for index, (mine_row, theirs_row) in enumerate(view.row_pairs)
                    if mine_row is None
                    and theirs_row is not None
                    and app.ws_b_val(_SHEET).cell(
                        row=theirs_row, column=2,
                    ).value == key
                )
                assert view._copy_selected_row(
                    "B2A", override_pair_idx=pair_index,
                ), key

            output = app.build_manual_merge_output_file()
            workbook = load_workbook(output, data_only=False)
            try:
                sheet = workbook[_SHEET]
                saved_keys = [
                    sheet.cell(row=row, column=2).value
                    for row in range(5, sheet.max_row + 1)
                ]
            finally:
                workbook.close()
            assert saved_keys == [
                common[0][1],
                None,
                *(row[1] for row in mine_additions),
                *remote_keys,
            ], saved_keys
        finally:
            if app is not None:
                app._shutdown_root()
            sm.SowMergeApp._schedule_formula_cache_prompt = original_scheduler
            if output is not None:
                Path(output).unlink(missing_ok=True)

    print("SMOKE_THREE_WAY_INDEPENDENT_DECLARED_APPENDS_OK")


if __name__ == "__main__":
    main()
