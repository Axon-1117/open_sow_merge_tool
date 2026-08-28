"""Regressions for Condition region history and virtual tail visibility."""

from __future__ import annotations

import os
import tempfile
import time

from openpyxl import Workbook

import sow_merge_tool as sm


_SHEET = "ConditionData@design"
_HEADER = [
    [
        "id@rely_id#key#ConditionData",
        "key@index_uniq_cached@constid",
        "type@ref_Condition",
        "param1",
    ],
    ["uint32", "string", "@ref", "string"],
    ["generated id", "condition key", "condition type", "parameter"],
    [None, None, None, None],
]


def _write(path: str, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _SHEET
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _pump(app) -> None:
    app.root.update_idletasks()
    app.root.update()


def _wait(app, predicate, label: str, timeout: float = 30.0) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        _pump(app)
        if predicate():
            return
        time.sleep(0.01)
    raise AssertionError(f"timeout waiting for {label}")


def _wait_ready(app, *, editable: bool):
    def ready():
        if editable:
            app._request_edit_preload()
        view = app.sheet_views.get(_SHEET)
        return bool(
            view is not None
            and view._data_ready
            and view._prepared_complete
            and app._is_sheet_exact_current(_SHEET)
            and (not editable or app._edit_workbooks_ready())
            and (not editable or view._derive_lifecycle_state() == "READY")
        )

    _wait(app, ready, "exact Condition surface")
    return app.sheet_views[_SHEET]


def _sheet_state(app, side: str):
    edit = app.ws_a_edit(_SHEET) if side == "A" else app.ws_b_edit(_SHEET)
    value = app.ws_a_val(_SHEET) if side == "A" else app.ws_b_val(_SHEET)
    rows = max(int(edit.max_row or 1), int(value.max_row or 1))
    cols = max(int(edit.max_column or 1), int(value.max_column or 1))
    return tuple(
        tuple(
            (
                edit.cell(row=row, column=col).value,
                value.cell(row=row, column=col).value,
            )
            for col in range(1, cols + 1)
        )
        for row in range(1, rows + 1)
    )


def _model_state(view):
    return (
        tuple(view.row_pairs),
        tuple(sorted((int(key), frozenset(value)) for key, value in view.pair_diff_cols.items())),
        tuple(sorted((int(key), frozenset(value)) for key, value in view.pair_base_diff_cols.items())),
        tuple(sorted(view.mine_to_base_row.items())),
        tuple(sorted(view.theirs_to_base_row.items())),
        tuple(sorted(view.pair_base_row_override.items())),
    )


def _assert_three_way_region_history(root: str) -> None:
    base_path = os.path.join(root, "history_base.xlsx")
    mine_path = os.path.join(root, "history_mine.xlsx")
    theirs_path = os.path.join(root, "history_theirs.xlsx")
    common = [[None, "quest_complete_402001", "quest_complete", 402001]]
    mine_additions = [
        [None, "grid_discover_day_2", "server_open_day", 2],
        [None, "grid_discover_day_3", "server_open_day", 3],
    ]
    theirs_additions = [
        [None, f"recharge_{condition}", "recharge", condition]
        for condition in range(30102, 30107)
    ]
    _write(base_path, _HEADER + common)
    _write(mine_path, _HEADER + common + [[None, None, None, None]] + mine_additions)
    _write(theirs_path, _HEADER + common + theirs_additions)

    app = None
    try:
        app = sm.SowMergeApp(
            mine_path,
            theirs_path,
            merge_mode=True,
            merged_path=os.path.join(root, "history_merged.xlsx"),
            base_path=base_path,
            initial_sheet=_SHEET,
        )
        view = _wait_ready(app, editable=True)
        before_sheet = _sheet_state(app, "A")
        before_model = _model_state(view)
        remote_pair = next(
            pair_idx
            for pair_idx, (mine_row, theirs_row) in enumerate(view.row_pairs)
            if mine_row is None and theirs_row is not None
        )
        view.selected_pair_idx = remote_pair
        view._last_selected_line = view.row_to_line.get(remote_pair, 1)
        view.selected_excel_row_a = None
        view.selected_excel_row_b = view.row_pairs[remote_pair][1]
        view._copy_selected_region("B2A")
        _pump(app)

        assert app.undo_stack and app.undo_stack[-1].get("kind") == "compound"
        compound = app.undo_stack[-1]
        assert len(compound.get("actions") or ()) >= 2, compound
        after_sheet = _sheet_state(app, "A")
        after_model = _model_state(view)
        assert after_sheet != before_sheet

        view._undo_last_action()
        _pump(app)
        assert _sheet_state(app, "A") == before_sheet
        assert _model_state(view) == before_model
        assert len(app.redo_stack) == 1
        assert app.redo_stack[-1].get("kind") == "compound"

        view._redo_last_action()
        _pump(app)
        assert _sheet_state(app, "A") == after_sheet
        assert _model_state(view) == after_model
        assert not app.redo_stack
        assert app.undo_stack and app.undo_stack[-1].get("kind") == "compound"
    finally:
        if app is not None:
            app._shutdown_root()


def _assert_two_way_tail_is_visible(root: str) -> None:
    base_path = os.path.join(root, "tail_base.xlsx")
    mine_path = os.path.join(root, "tail_mine.xlsx")
    # Four schema rows + 6500 common records = physical row 6504.
    common = [
        [None, f"condition_{row}", "condition", row]
        for row in range(1, 6501)
    ]
    additions = [
        [None, "grid_discover_day_2", "server_open_day", 2],
        [None, "city_event_reward_3", "city_event_reward", 3],
        [None, "grid_discover_day_3", "server_open_day", 3],
    ]
    _write(base_path, _HEADER + common)
    _write(mine_path, _HEADER + common + additions)

    app = None
    try:
        app = sm.SowMergeApp(
            base_path,
            mine_path,
            merge_mode=False,
            initial_sheet=_SHEET,
        )
        view = _wait_ready(app, editable=False)
        if not bool(view.only_diff_var.get()):
            view.only_diff_var.set(1)
            view._toggle_only_diff()
            _wait(
                app,
                lambda: bool(view.only_diff_var.get()) and not view._mode_switch_pending,
                "only-difference mode",
            )
        assert view.row_pairs[view._full_display_rows[-1]][1] == 6507

        view.only_diff_var.set(0)
        view._toggle_only_diff()
        _wait(
            app,
            lambda: not bool(view.only_diff_var.get()) and not view._mode_switch_pending,
            "full mode",
        )
        view._publish_virtual_window(10**9)
        _pump(app)

        last_line = len(view.display_rows)
        last_pair = view.row_pairs[view.display_rows[-1]]
        first_info = view.right.dlineinfo("1.0")
        last_info = view.right.dlineinfo(f"{last_line}.0")
        assert last_pair[1] == 6507, (last_pair, view._virtual_window_start)
        assert first_info is not None and last_info is not None
        assert int(last_info[3]) == int(first_info[3]), (first_info, last_info)
        right_headers = view.right_ln.get("1.0", "end-1c").splitlines()
        assert right_headers[-1].strip() == "6507", right_headers[-5:]
    finally:
        if app is not None:
            app._shutdown_root()


def _assert_three_way_right_tail_and_alignment(root: str) -> None:
    base_path = os.path.join(root, "three_way_tail_base.xlsx")
    mine_path = os.path.join(root, "three_way_tail_mine.xlsx")
    theirs_path = os.path.join(root, "three_way_tail_theirs.xlsx")
    common = [
        [None, f"condition_{row}", "condition", row]
        for row in range(1, 6501)
    ]
    mine_additions = [
        [None, "grid_discover_day_2", "server_open_day", 2],
        [None, "grid_discover_day_3", "server_open_day", 3],
    ]
    theirs_additions = [
        [None, f"recharge_{condition}", "recharge", condition]
        for condition in range(30102, 30107)
    ]
    base_rows = _HEADER + common
    mine_rows = _HEADER + common + [[None, None, None, None]] + mine_additions
    theirs_rows = _HEADER + common + theirs_additions
    _write(base_path, base_rows)
    _write(mine_path, mine_rows)
    _write(theirs_path, theirs_rows)

    app = None
    try:
        app = sm.SowMergeApp(
            mine_path,
            theirs_path,
            merge_mode=True,
            merged_path=os.path.join(root, "three_way_tail_merged.xlsx"),
            base_path=base_path,
            initial_sheet=_SHEET,
        )
        view = _wait_ready(app, editable=False)
        expected_a_rows = set(range(1, len(mine_rows) + 1))
        expected_b_rows = set(range(1, len(theirs_rows) + 1))
        actual_a_rows = [row_a for row_a, _row_b in view.row_pairs if row_a is not None]
        actual_b_rows = [row_b for _row_a, row_b in view.row_pairs if row_b is not None]
        assert len(actual_a_rows) == len(set(actual_a_rows))
        assert len(actual_b_rows) == len(set(actual_b_rows))
        assert set(actual_a_rows) == expected_a_rows
        assert set(actual_b_rows) == expected_b_rows

        for row_a, row_b in view.row_pairs:
            if row_a is None or row_b is None or row_a <= len(_HEADER) or row_b <= len(_HEADER):
                continue
            key_a = mine_rows[row_a - 1][1]
            key_b = theirs_rows[row_b - 1][1]
            if key_a is not None and key_b is not None:
                assert key_a == key_b, (row_a, row_b, key_a, key_b)

        right_tail_row = len(theirs_rows)
        expected_tail = (
            (len(base_rows) + 1, None),
            (len(base_rows) + 2, None),
            (len(base_rows) + 3, None),
            *((None, row) for row in range(len(base_rows) + 1, right_tail_row + 1)),
        )
        assert tuple(view.row_pairs[-len(expected_tail):]) == expected_tail
        assert all(
            set(view.pair_diff_cols.get(pair_idx, ())) == {-1}
            for pair_idx in range(len(view.row_pairs) - len(expected_tail), len(view.row_pairs))
        )
        assert view.row_pairs[view._full_display_rows[-1]][1] == right_tail_row
        if bool(view.only_diff_var.get()):
            view.only_diff_var.set(0)
            view._toggle_only_diff()
            _wait(
                app,
                lambda: not bool(view.only_diff_var.get()) and not view._mode_switch_pending,
                "three-way full mode",
            )
        view._publish_virtual_window(10**9)
        _pump(app)

        last_line = len(view.display_rows)
        last_pair = view.row_pairs[view.display_rows[-1]]
        first_info = view.right.dlineinfo("1.0")
        last_info = view.right.dlineinfo(f"{last_line}.0")
        assert last_pair[1] == right_tail_row, (last_pair, view._virtual_window_start)
        assert first_info is not None and last_info is not None
        assert int(last_info[3]) == int(first_info[3]), (first_info, last_info)
        right_headers = view.right_ln.get("1.0", "end-1c").splitlines()
        assert right_headers[-1].strip() == str(right_tail_row), right_headers[-5:]
    finally:
        if app is not None:
            app._shutdown_root()


def main() -> None:
    original_prompt = sm.SowMergeApp._schedule_formula_cache_prompt
    sm.SowMergeApp._schedule_formula_cache_prompt = lambda _self: None
    try:
        with tempfile.TemporaryDirectory(prefix="sow_condition_history_tail_") as root:
            _assert_three_way_region_history(root)
            _assert_two_way_tail_is_visible(root)
            _assert_three_way_right_tail_and_alignment(root)
    finally:
        sm.SowMergeApp._schedule_formula_cache_prompt = original_prompt
    print("SMOKE_CONDITION_HISTORY_VISIBLE_TAIL_OK")


if __name__ == "__main__":
    main()
