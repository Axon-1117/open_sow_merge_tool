"""Persist executed physical operations independently of comparison readiness."""
import os
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import sow_merge_tool as sm
from _smoke_test_column_native_save_replay import _fake_app


def book(path, missing=False):
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.append(["id", "value"])
    wb.active.append([1, "original"])
    if not missing:
        ws = wb.create_sheet("ResidentLink@design")
        ws.append(["id", "value"])
        ws.append([7, "mine-only"])
    wb.save(path)
    wb.close()


def forbidden(*args, **kwargs):
    raise AssertionError("saving must not read or rebuild comparison views")


def main():
    with tempfile.TemporaryDirectory(prefix="sow-save-failed-") as root:
        mine, theirs = [os.path.join(root, name + ".xlsx") for name in ("mine", "theirs")]
        book(mine)
        book(theirs, missing=True)
        original = open(mine, "rb").read()
        app = _fake_app(mine, theirs)
        app.has_base = False
        app._wb_a_edit = load_workbook(mine)
        app._wb_b_edit = load_workbook(theirs)
        app._wb_a_val = load_workbook(mine)
        app._wb_b_val = load_workbook(theirs)
        app._wb_a_edit["Data"]["B2"] = "executed edit"
        app._wb_a_val["Data"]["B2"] = "executed edit"
        app.manual_a_cell_ops[("Data", 2, 2)] = "executed edit"
        app.modified_a = True
        app.selected_sheet = "ResidentLink@design"
        app._ensure_live_column_mappings_current = forbidden
        app._show_exact_readiness_modal = forbidden
        try:
            for state in ("FAILED", "CALCULATING", "UNRESOLVED", "CANCELED", "STALE", "UNLOADED"):
                app.sheet_views = {} if state == "UNLOADED" else {
                    app.selected_sheet: SimpleNamespace(_data_ready=True,
                        _derive_lifecycle_state=lambda: state,
                        _ensure_column_projection_current=forbidden)}
                assert app._guard_save_readiness("保存 Merged", "A"), state
                output = app.build_manual_merge_output_file()
                try:
                    result = load_workbook(output)
                    assert result["Data"]["B2"].value == "executed edit", state
                    assert result["ResidentLink@design"]["B2"].value == "mine-only", state
                    result.close()
                finally:
                    os.remove(output)
                assert open(mine, "rb").read() == original
                assert app.manual_a_cell_ops[("Data", 2, 2)] == "executed edit"

            # Real public save entry point writes local mine and exits only after success.
            app.merged_path = mine
            app.merge_mode = True
            app.has_base = True
            app._wb_base_edit = load_workbook(theirs)
            app._wb_base_val = load_workbook(theirs)
            app.initial_conflict_cell_count = 0
            app.sheet_level_conflicts = []
            app._begin_interactive_action = lambda: None
            app._end_interactive_action = lambda: None
            app._with_progress = lambda title, text, task, **kwargs: task(lambda *args: None)
            app._shutdown_root = lambda: None
            with patch.object(sm.messagebox, "askyesno", return_value=True), \
                 patch.object(sm.messagebox, "showinfo"), \
                 patch.object(sm.messagebox, "showerror", side_effect=forbidden), \
                 patch.object(sm, "_has_svn_conflict_artifacts", return_value=False):
                try:
                    app.save_merged_and_exit()
                except SystemExit as exc:
                    assert exc.code == 0
                else:
                    raise AssertionError("public merged save did not complete")
            saved = load_workbook(mine)
            assert saved["Data"]["B2"].value == "executed edit"
            assert saved["ResidentLink@design"]["B2"].value == "mine-only"
            saved.close()
            assert not app.modified_a
        finally:
            for attr in ("_wb_a_edit", "_wb_b_edit", "_wb_a_val", "_wb_b_val", "_wb_base_edit", "_wb_base_val"):
                wb = getattr(app, attr, None)
                if wb is not None:
                    wb.close()
    print("PASS: failed/pending/unresolved/stale/unopened comparison saves executed edits to mine")


if __name__ == "__main__":
    main()
